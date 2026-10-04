"""Importación de catálogos de proveedores (Truper u otros) desde Excel o CSV.

Flujo:
  1. leer_archivo()       -> convierte el archivo en una lista de filas
  2. detectar_encabezado() / adivinar_columnas() -> propone qué columna es qué
  3. aplicar_importacion() -> valida y guarda los productos en una transacción
"""

import csv
import io
import unicodedata

from .utils import a_centavos

CAMPOS = {
    "codigo": "Código",
    "nombre": "Descripción / nombre",
    "precio": "Precio",
    "stock": "Existencias (opcional)",
    "stock_minimo": "Stock mínimo (opcional)",
}
OBLIGATORIOS = ("codigo", "nombre", "precio")

# Palabras con las que suelen venir los encabezados en las listas de precios.
PALABRAS_CLAVE = {
    "codigo": ("codigo", "clave", "sku", "cod", "articulo", "no. parte", "modelo", "upc", "ean"),
    "nombre": ("descripcion", "nombre", "producto"),
    "precio": ("precio publico", "publico", "precio venta", "precio", "p.v.", "pvp", "importe", "costo"),
    # stock_minimo va antes que stock para que "Stock mínimo" no se confunda con "Stock".
    "stock_minimo": ("minimo", "minima", "stock min"),
    "stock": ("existencia", "stock", "inventario"),
}


class ErrorImportacion(Exception):
    pass


def _normalizar(texto):
    """'Descripción ' -> 'descripcion' (sin acentos, minúsculas, sin espacios extra)."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def _celda_a_texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # 10250.0 -> "10250"
    return str(valor).strip()


def leer_archivo(nombre_archivo, contenido):
    """Devuelve una lista de filas (listas de textos) a partir de un .xlsx o .csv."""
    nombre = nombre_archivo.lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        return _leer_excel(contenido)
    if nombre.endswith((".csv", ".txt")):
        return _leer_csv(contenido)
    if nombre.endswith(".xls"):
        raise ErrorImportacion(
            "Los archivos .xls (Excel antiguo) no se pueden leer. Ábrelo en Excel y usa "
            "Archivo > Guardar como > Libro de Excel (.xlsx)."
        )
    if nombre.endswith(".pdf"):
        raise ErrorImportacion(
            "Los catálogos en PDF no se pueden importar directamente. Pide a tu distribuidor "
            "la lista de precios en Excel."
        )
    raise ErrorImportacion("Sube un archivo de Excel (.xlsx) o CSV (.csv).")


def _leer_excel(contenido):
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorImportacion("No se pudo abrir el archivo de Excel. Revisa que no esté dañado.")

    # Usa la hoja con más filas (las listas de precios a veces traen portada).
    mejor = []
    for hoja in libro.worksheets:
        filas = [[_celda_a_texto(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]
        filas = _quitar_vacias_al_final(filas)
        if sum(1 for f in filas if any(f)) > sum(1 for f in mejor if any(f)):
            mejor = filas
    libro.close()
    if not mejor:
        raise ErrorImportacion("El archivo de Excel está vacío.")
    return mejor


def _leer_csv(contenido):
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:5000], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), dialecto)]
    filas = _quitar_vacias_al_final(filas)
    if not filas:
        raise ErrorImportacion("El archivo CSV está vacío.")
    return filas


def _quitar_vacias_al_final(filas):
    while filas and not any(filas[-1]):
        filas.pop()
    return filas


def detectar_encabezado(filas):
    """Busca en las primeras 25 filas la que parece ser la de encabezados."""
    mejor_indice, mejor_puntos = 0, 0
    for i, fila in enumerate(filas[:25]):
        puntos = 0
        for celda in fila:
            c = _normalizar(celda)
            if c and any(p in c for palabras in PALABRAS_CLAVE.values() for p in palabras):
                puntos += 1
        if puntos > mejor_puntos:
            mejor_indice, mejor_puntos = i, puntos
    return mejor_indice


def adivinar_columnas(encabezados):
    """Propone qué número de columna corresponde a cada campo."""
    normalizados = [_normalizar(e) for e in encabezados]
    usadas = set()
    resultado = {}
    for campo, palabras in PALABRAS_CLAVE.items():
        for palabra in palabras:  # en orden de preferencia
            for i, enc in enumerate(normalizados):
                if i not in usadas and enc and palabra in enc:
                    resultado[campo] = i
                    usadas.add(i)
                    break
            if campo in resultado:
                break
    return resultado


def aplicar_importacion(db, filas, columnas, si_existe="actualizar", aumento_pct=0, fila_inicial=1):
    """Guarda los productos. Devuelve un resumen con nuevos, actualizados, omitidos y errores.

    filas: filas de datos (sin la fila de encabezados)
    columnas: {"codigo": 0, "nombre": 2, "precio": 5, "stock": None, ...}
    si_existe: "actualizar" (precio y nombre) u "omitir"
    aumento_pct: porcentaje a sumar al precio del archivo (por ejemplo, tu margen de ganancia)
    fila_inicial: número de fila en el archivo original de la primera fila de datos,
                  para que los mensajes de error coincidan con lo que se ve en Excel
    """
    for campo in OBLIGATORIOS:
        if columnas.get(campo) is None:
            raise ErrorImportacion(f"Elige qué columna corresponde a \"{CAMPOS[campo]}\".")

    def celda(fila, campo):
        i = columnas.get(campo)
        return fila[i].strip() if i is not None and i < len(fila) else ""

    def entero(texto, campo, n):
        if not texto:
            return None
        try:
            valor = int(float(texto.replace(",", "")))
            if valor < 0:
                raise ValueError
            return valor
        except ValueError:
            raise ValueError(f"fila {n}: {CAMPOS[campo].split(' (')[0].lower()} \"{texto}\" no es un número válido")

    factor = 1 + (aumento_pct or 0) / 100
    resumen = {"nuevos": 0, "actualizados": 0, "omitidos": 0, "errores": []}
    vistos = set()

    try:
        db.execute("BEGIN")
        existentes = {
            r["codigo"]: r for r in db.execute("SELECT id, codigo, activo FROM productos").fetchall()
        }
        for n, fila in enumerate(filas, start=fila_inicial):
            codigo = celda(fila, "codigo")
            nombre = celda(fila, "nombre")
            texto_precio = celda(fila, "precio")

            if not codigo and not nombre:
                continue  # fila vacía o separador
            try:
                if not codigo:
                    raise ValueError(f"fila {n}: falta el código")
                if not nombre:
                    raise ValueError(f"fila {n}: falta la descripción")
                try:
                    precio = round(a_centavos(texto_precio) * factor)
                except ValueError:
                    raise ValueError(f"fila {n}: el precio \"{texto_precio}\" no es válido")
                stock = entero(celda(fila, "stock"), "stock", n)
                stock_minimo = entero(celda(fila, "stock_minimo"), "stock_minimo", n)
            except ValueError as e:
                resumen["omitidos"] += 1
                resumen["errores"].append(str(e))
                continue

            if codigo in vistos:
                resumen["omitidos"] += 1
                resumen["errores"].append(f"fila {n}: el código {codigo} está repetido en el archivo")
                continue
            vistos.add(codigo)

            actual = existentes.get(codigo)
            if actual is None:
                db.execute(
                    "INSERT INTO productos (codigo, nombre, precio_centavos, stock, stock_minimo)"
                    " VALUES (?, ?, ?, ?, ?)",
                    (codigo, nombre, precio, stock or 0, stock_minimo or 0),
                )
                resumen["nuevos"] += 1
            elif si_existe == "omitir":
                resumen["omitidos"] += 1
            else:
                cambios = {"nombre": nombre, "precio_centavos": precio, "activo": 1}
                if stock is not None:
                    cambios["stock"] = stock
                if stock_minimo is not None:
                    cambios["stock_minimo"] = stock_minimo
                asignaciones = ", ".join(f"{k} = :{k}" for k in cambios)
                db.execute(
                    f"UPDATE productos SET {asignaciones} WHERE id = :id",
                    {**cambios, "id": actual["id"]},
                )
                resumen["actualizados"] += 1

        db.commit()
    except Exception:
        db.rollback()
        raise
    return resumen
